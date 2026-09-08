from __future__ import annotations

import argparse
import csv
import datetime
import html
import json
import re
import sys
import time
import urllib.error
import urllib.parse
import urllib.request
from dataclasses import dataclass, field
from html.parser import HTMLParser
from pathlib import Path
from typing import Any

from docx import Document
from openpyxl import Workbook


BASE_URL = "http://ay-software.in.th/ptn/"

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
if hasattr(sys.stderr, "reconfigure"):
    sys.stderr.reconfigure(encoding="utf-8", errors="replace")


class FormParser(HTMLParser):
    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.forms: list[dict[str, Any]] = []
        self._form: dict[str, Any] | None = None
        self._button_text: list[str] | None = None
        self._textarea_name: str | None = None
        self._textarea_chunks: list[str] = []

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        attrs_d = {k.lower(): (v or "") for k, v in attrs}
        if tag == "form":
            self._form = {
                "action": attrs_d.get("action", ""),
                "method": attrs_d.get("method", "get").lower(),
                "inputs": {},
                "buttons": [],
                "textareas": {},
            }
        elif self._form is not None and tag == "input":
            name = attrs_d.get("name")
            if name:
                self._form["inputs"][name] = attrs_d.get("value", "")
        elif self._form is not None and tag == "textarea":
            self._textarea_name = attrs_d.get("name")
            self._textarea_chunks = []
        elif self._form is not None and tag == "button":
            self._button_text = []

    def handle_data(self, data: str) -> None:
        if self._button_text is not None:
            self._button_text.append(data)
        if self._textarea_name is not None:
            self._textarea_chunks.append(data)

    def handle_endtag(self, tag: str) -> None:
        if tag == "button" and self._form is not None and self._button_text is not None:
            self._form["buttons"].append(clean_text(" ".join(self._button_text)))
            self._button_text = None
        elif tag == "textarea" and self._form is not None and self._textarea_name is not None:
            self._form["textareas"][self._textarea_name] = html.unescape("".join(self._textarea_chunks)).strip()
            self._textarea_name = None
            self._textarea_chunks = []
        elif tag == "form" and self._form is not None:
            self.forms.append(self._form)
            self._form = None


class TextParser(HTMLParser):
    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.parts: list[str] = []

    def handle_data(self, data: str) -> None:
        if data.strip():
            self.parts.append(data.strip())

    @property
    def text(self) -> str:
        return clean_text(" ".join(self.parts))


def clean_text(value: str) -> str:
    value = html.unescape(value or "")
    value = re.sub(r"\s+", " ", value, flags=re.MULTILINE).strip()
    return value


def parse_forms(source: str) -> list[dict[str, Any]]:
    parser = FormParser()
    parser.feed(source)
    return parser.forms


def html_text(source: str) -> str:
    parser = TextParser()
    parser.feed(source)
    return parser.text


def normalize_filename(value: str) -> str:
    value = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", value)
    value = re.sub(r"\s+", " ", value).strip()
    return value[:120] or "unknown"


@dataclass
class SubjectRecord:
    teacher_username: str
    teacher_id: str
    teacher_name: str
    level: str
    class_code: str
    subject_code: str
    subject_name: str
    teaching_id: str
    credit: str = ""
    course_description: str = ""
    learning_outcomes: str = ""
    score_before: str = ""
    score_mid: str = ""
    score_after: str = ""
    score_final: str = ""
    score_ratio: str = ""
    attendance_checked: int = 0
    attendance_expected: int | None = None
    score_items_count: int = 0
    score_items_total: float = 0.0
    notes: list[str] = field(default_factory=list)

    @property
    def class_for_pt05(self) -> str:
        if self.class_code.isdigit():
            return self.class_code
        if "/" in self.class_code:
            return self.class_code.rsplit("/", 1)[-1]
        return "รวมชั้น" if self.class_code else ""

    @property
    def course_status(self) -> str:
        return "มี" if self.course_description else "ไม่มี"

    @property
    def outcome_status(self) -> str:
        return "มี" if self.learning_outcomes else "ไม่มี"

    @property
    def attendance_status(self) -> str:
        if self.attendance_expected:
            return f"{self.attendance_checked}/{self.attendance_expected}"
        return str(self.attendance_checked) if self.attendance_checked else ""

    @property
    def score_status(self) -> str:
        if self.score_items_total:
            total = int(self.score_items_total) if self.score_items_total.is_integer() else self.score_items_total
            return f"{total}/100"
        return ""

    @property
    def notes_text(self) -> str:
        return "; ".join(dict.fromkeys([n for n in self.notes if n]))


class SssClient:
    def __init__(self, base_url: str = BASE_URL, timeout: int = 30, delay: float = 0.2) -> None:
        self.base_url = base_url.rstrip("/") + "/"
        self.timeout = timeout
        self.delay = delay
        self.opener = urllib.request.build_opener(urllib.request.HTTPCookieProcessor())
        self.opener.addheaders = [
            ("User-Agent", "Mozilla/5.0 PT05Tool/1.0"),
            ("Accept", "text/html,application/xhtml+xml,application/xml;q=0.9,*/*;q=0.8"),
        ]

    def post(self, endpoint: str, data: dict[str, str]) -> str:
        url = urllib.parse.urljoin(self.base_url, endpoint)
        encoded = urllib.parse.urlencode(data).encode("utf-8")
        req = urllib.request.Request(url, data=encoded, method="POST")
        req.add_header("Content-Type", "application/x-www-form-urlencoded; charset=UTF-8")
        try:
            with self.opener.open(req, timeout=self.timeout) as resp:
                raw = resp.read()
        except urllib.error.URLError as exc:
            raise RuntimeError(f"เปิด {endpoint} ไม่สำเร็จ: {exc}") from exc
        time.sleep(self.delay)
        return raw.decode("utf-8", errors="replace")

    def login(self, username: str, password: str) -> dict[str, str]:
        html1 = self.post("chk_user.php", {"user": username, "pass": password, "menu": "3"})
        forms = parse_forms(html1)
        confirm = next((f for f in forms if f.get("action") == "index_teacher.php"), None)
        if not confirm:
            raise RuntimeError(f"ล็อกอิน {username} ไม่สำเร็จ หรือไม่ใช่บัญชีครู")
        data = {k: str(v) for k, v in confirm["inputs"].items()}
        self.post("index_teacher.php", data)
        return data


def load_teachers(path: Path) -> list[dict[str, str]]:
    if path.suffix.lower() in {".xlsx", ".xlsm"}:
        from openpyxl import load_workbook

        wb = load_workbook(path, data_only=True)
        ws = wb.active
        headers = [str(c.value or "").strip().lower() for c in ws[1]]
        rows = []
        for row in ws.iter_rows(min_row=2, values_only=True):
            item = {headers[i]: str(value or "").strip() for i, value in enumerate(row) if i < len(headers)}
            if item.get("username") and item.get("password"):
                rows.append(item)
        return rows

    with path.open("r", encoding="utf-8-sig", newline="") as f:
        rows = []
        for row in csv.DictReader(f):
            item = {str(k).strip().lower(): str(v or "").strip() for k, v in row.items()}
            if item.get("username") and item.get("password"):
                rows.append(item)
        return rows


def extract_subjects(list_html: str, teacher_username: str, teacher_info: dict[str, str]) -> list[SubjectRecord]:
    forms = parse_forms(list_html)
    records: dict[tuple[str, str, str, str], SubjectRecord] = {}
    for form in forms:
        inputs = form["inputs"]
        action = form.get("action", "")
        if action not in {"del_test.php", "re_teaching.php", "sco_confix.php", "sco_confix2.php"}:
            continue
        required = ["id_sub", "level_t", "class_t", "id_teaching"]
        if not all(inputs.get(k) for k in required):
            continue

        key = (inputs["id_sub"], inputs["level_t"], inputs["class_t"], inputs["id_teaching"])
        record = records.get(key)
        if not record:
            record = SubjectRecord(
                teacher_username=teacher_username,
                teacher_id=teacher_info.get("id_1", ""),
                teacher_name=clean_text(f"{teacher_info.get('name_1', '')} {teacher_info.get('lname_1', '')}"),
                subject_code=inputs["id_sub"],
                subject_name=inputs.get("name_sub", ""),
                level=inputs["level_t"],
                class_code=inputs["class_t"],
                teaching_id=inputs["id_teaching"],
            )
            records[key] = record
        if inputs.get("name_sub"):
            record.subject_name = inputs["name_sub"]
        if not record.subject_name:
            text = " ".join(form.get("buttons", []))
            match = re.search(rf"{re.escape(record.subject_code)}\s+(.+?)\s*\(", text)
            if match:
                record.subject_name = clean_text(match.group(1))
    return list(records.values())


def extract_course_info(source: str, record: SubjectRecord) -> None:
    forms = parse_forms(source)
    for form in forms:
        areas = form.get("textareas", {})
        if "cou1" in areas:
            record.course_description = areas.get("cou1", "")
        if "cou3" in areas:
            record.learning_outcomes = areas.get("cou3", "")
    text = html_text(source)
    credit = re.search(r"จำนวน\s*([0-9.]+)\s*หน่วยกิต", text)
    if credit:
        record.credit = credit.group(1)
        record.attendance_expected = expected_periods(record.credit)
    if not record.course_description:
        record.notes.append("ไม่มีคำอธิบายรายวิชา")
    if not record.learning_outcomes:
        record.notes.append("ไม่มีตัวชี้วัด/ผลการเรียนรู้")


def extract_score_config(source: str, record: SubjectRecord) -> None:
    forms = parse_forms(source)
    form = next((f for f in forms if f.get("action") == "sco_confix.php"), None)
    if not form:
        record.notes.append("ไม่พบหน้ากำหนดคะแนน")
        return
    inputs = form["inputs"]
    record.score_before = inputs.get("sco_before", "")
    record.score_mid = inputs.get("sco_mid", "")
    record.score_after = inputs.get("sco_after", "")
    record.score_final = inputs.get("sco_final", "")
    during = sum_ints(record.score_before, record.score_mid, record.score_after)
    final = to_int(record.score_final)
    if during is not None and final is not None:
        record.score_ratio = f"{during}:{final}"


def extract_attendance(source: str, record: SubjectRecord) -> None:
    dates = re.findall(r"\b\d{2}/\d{2}/25\d{2}\b", source)
    record.attendance_checked = len(dates)
    if record.attendance_checked == 0:
        record.notes.append("ยังไม่พบการเช็คเวลาเรียน")


def extract_score_items(source: str, record: SubjectRecord) -> None:
    scores = [float(v.replace(",", "")) for v in re.findall(r"ฐานคะแนน:\s*([0-9]+(?:\.[0-9]+)?)", source)]
    record.score_items_count = len(scores)
    record.score_items_total = sum(scores)
    if record.score_items_count == 0:
        record.notes.append("ยังไม่พบรายการคะแนน")
    elif record.score_items_total > 100:
        total = int(record.score_items_total) if record.score_items_total.is_integer() else record.score_items_total
        record.notes.append(f"รายการคะแนนรวม {total} เกิน 100")


def expected_periods(credit: str) -> int | None:
    try:
        value = float(credit)
    except ValueError:
        return None
    expected = int(round(value * 40))
    return expected if expected > 0 else None


def to_int(value: str) -> int | None:
    try:
        return int(float(str(value).strip()))
    except ValueError:
        return None


def sum_ints(*values: str) -> int | None:
    ints = [to_int(v) for v in values]
    if any(v is None for v in ints):
        return None
    return sum(v for v in ints if v is not None)


def fetch_teacher(username: str, password: str, raw_dir: Path | None = None) -> list[SubjectRecord]:
    client = SssClient()
    info = client.login(username, password)
    list_html = client.post("re_list_teaching.php", {"id_1": info.get("id_1", "")})
    save_raw(raw_dir, username, "re_list_teaching", list_html)
    subjects = extract_subjects(list_html, username, info)

    for record in subjects:
        base = {
            "id_1": record.teacher_id,
            "id_sub": record.subject_code,
            "level_t": record.level,
            "class_t": record.class_code,
            "id_teaching": record.teaching_id,
        }
        course_html = client.post("sco_confix2.php", base)
        save_raw(raw_dir, username, f"{record.teaching_id}_sco_confix2", course_html)
        extract_course_info(course_html, record)

        config_html = client.post("sco_confix.php", base)
        save_raw(raw_dir, username, f"{record.teaching_id}_sco_confix", config_html)
        extract_score_config(config_html, record)

        attendance_html = client.post("re_teaching.php", base)
        save_raw(raw_dir, username, f"{record.teaching_id}_re_teaching", attendance_html)
        extract_attendance(attendance_html, record)

        score_payload = dict(base)
        score_payload["name_sub"] = record.subject_name
        score_html = client.post("del_test.php", score_payload)
        save_raw(raw_dir, username, f"{record.teaching_id}_del_test", score_html)
        extract_score_items(score_html, record)

    return subjects


def save_raw(raw_dir: Path | None, username: str, name: str, content: str) -> None:
    if raw_dir is None:
        return
    folder = raw_dir / normalize_filename(username)
    folder.mkdir(parents=True, exist_ok=True)
    (folder / f"{normalize_filename(name)}.html").write_text(content, encoding="utf-8")


def write_summary(records: list[SubjectRecord], output: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "PT05"
    headers = [
        "username",
        "teacher_id",
        "teacher_name",
        "level",
        "class",
        "subject_code",
        "subject_name",
        "course_description_status",
        "learning_outcomes_status",
        "attendance_checked",
        "attendance_expected",
        "score_ratio",
        "score_items_count",
        "score_items_total",
        "notes",
    ]
    ws.append(headers)
    for r in records:
        ws.append([
            r.teacher_username,
            r.teacher_id,
            r.teacher_name,
            r.level,
            r.class_code,
            r.subject_code,
            r.subject_name,
            r.course_status,
            r.outcome_status,
            r.attendance_checked,
            r.attendance_expected or "",
            r.score_ratio,
            r.score_items_count,
            r.score_items_total,
            r.notes_text,
        ])
    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        ws.column_dimensions[col[0].column_letter].width = min(max(max_len + 2, 10), 40)
    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output)

THAI_MONTHS = [
    "", "มกราคม", "กุมภาพันธ์", "มีนาคม", "เมษายน", "พฤษภาคม",
    "มิถุนายน", "กรกฎาคม", "สิงหาคม", "กันยายน", "ตุลาคม", "พฤศจิกายน", "ธันวาคม"
]

THAI_MONTH_ABBR = {
    "ม.ค.": "มกราคม", "ม.ค": "มกราคม",
    "ก.พ.": "กุมภาพันธ์", "ก.พ": "กุมภาพันธ์",
    "มี.ค.": "มีนาคม", "มี.ค": "มีนาคม",
    "เม.ย.": "เมษายน", "เม.ย": "เมษายน",
    "พ.ค.": "พฤษภาคม", "พ.ค": "พฤษภาคม",
    "มิ.ย.": "มิถุนายน", "มิ.ย": "มิถุนายน",
    "ก.ค.": "กรกฎาคม", "ก.ค": "กรกฎาคม",
    "ส.ค.": "สิงหาคม", "ส.ค": "สิงหาคม",
    "ก.ย.": "กันยายน", "ก.ย": "กันยายน",
    "ต.ค.": "ตุลาคม", "ต.ค": "ตุลาคม",
    "พ.ย.": "พฤศจิกายน", "พ.ย": "พฤศจิกายน",
    "ธ.ค.": "ธันวาคม", "ธ.ค": "ธันวาคม"
}


def get_thai_date(date_str: str | None = None) -> tuple[str, str, str]:
    if not date_str or not date_str.strip():
        dt = datetime.datetime.now()
        return str(dt.day), THAI_MONTHS[dt.month], str(dt.year + 543)

    date_str = date_str.strip()

    # 1. Try ISO format: YYYY-MM-DD
    try:
        dt = datetime.datetime.strptime(date_str, "%Y-%m-%d")
        return str(dt.day), THAI_MONTHS[dt.month], str(dt.year + 543)
    except ValueError:
        pass

    # 2. Try DD/MM/YYYY or DD-MM-YYYY
    for fmt in ("%d/%m/%Y", "%d-%m-%Y"):
        try:
            dt = datetime.datetime.strptime(date_str, fmt)
            year = dt.year
            if year < 2400:
                year += 543
            return str(dt.day), THAI_MONTHS[dt.month], str(year)
        except ValueError:
            pass

    # 3. Try custom string like "4 กรกฎาคม 2569" or "4 ก.ค. 2569"
    parts = re.split(r"[\s\-/]+", date_str)
    if len(parts) == 3:
        day = parts[0]
        month_input = parts[1]
        year_input = parts[2]
        
        month_name = THAI_MONTH_ABBR.get(month_input, month_input)
        if month_name == month_input:
            for m in THAI_MONTHS:
                if m and (month_input in m or m in month_input):
                    month_name = m
                    break
        
        try:
            yr = int(year_input)
            if yr < 100:
                yr += 2500
            elif yr < 2400:
                yr += 543
            year_be = str(yr)
        except ValueError:
            year_be = year_input
            
        return day, month_name, year_be

    # Fallback to today
    dt = datetime.datetime.now()
    return str(dt.day), THAI_MONTHS[dt.month], str(dt.year + 543)


def fill_docx(template: Path, records: list[SubjectRecord], output: Path, term: str, year: str, date_str: str | None = None) -> None:
    doc = Document(template)
    teacher_name = records[0].teacher_name if records else ""
    day, month_name, year_be = get_thai_date(date_str)
    for paragraph in doc.paragraphs:
        text = paragraph.text
        if "ภาคเรียนที่" in text and "ครูประจำวิชา" in text:
            paragraph.text = (
                f"ภาคเรียนที่ {term}/{year} ครูประจำวิชา {teacher_name} "
                f"ตรวจสอบวันที่ {day} เดือน {month_name} พ.ศ. {year_be}"
            )
            break

    if not doc.tables:
        raise RuntimeError("template ไม่มีตาราง ปถ.05")
    table = doc.tables[0]
    ensure_table_rows(table, len(records) + 1)

    for idx, row in enumerate(table.rows[1:], start=0):
        if idx >= len(records):
            clear_row(row)
            continue
        r = records[idx]
        values = [
            r.level,
            r.subject_code,
            r.subject_name,
            r.class_for_pt05,
            r.course_status,
            r.outcome_status,
            r.attendance_status,
            r.score_ratio,
            r.score_status,
            r.notes_text,
        ]
        for cell, value in zip(row.cells, values):
            cell.text = str(value or "")

    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)


def ensure_table_rows(table: Any, wanted: int) -> None:
    while len(table.rows) < wanted:
        table.add_row()


def clear_row(row: Any) -> None:
    for cell in row.cells:
        cell.text = ""


def write_json(records: list[SubjectRecord], output: Path) -> None:
    data = [r.__dict__ | {
        "class_for_pt05": r.class_for_pt05,
        "course_status": r.course_status,
        "outcome_status": r.outcome_status,
        "attendance_status": r.attendance_status,
        "score_status": r.score_status,
        "notes_text": r.notes_text,
    } for r in records]
    output.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")


def main() -> int:
    parser = argparse.ArgumentParser(description="ดึงข้อมูลจาก SSS แล้วสร้างเอกสาร ปถ.05")
    parser.add_argument("--teachers", required=True, type=Path, help="ไฟล์ CSV/XLSX มีคอลัมน์ username,password")
    parser.add_argument("--template", required=True, type=Path, help="ไฟล์ template ปถ.05 .docx")
    parser.add_argument("--out", default=Path("output-pt05"), type=Path, help="โฟลเดอร์ผลลัพธ์")
    parser.add_argument("--term", default="1", help="ภาคเรียน")
    parser.add_argument("--year", default="2569", help="ปีการศึกษา")
    parser.add_argument("--date", help="วันที่ตรวจสอบ (เช่น '2026-07-04' หรือ '4 กรกฎาคม 2569')")
    parser.add_argument("--save-raw", action="store_true", help="บันทึก HTML ต้นทางไว้ตรวจสอบ")
    args = parser.parse_args()

    teachers = load_teachers(args.teachers)
    if not teachers:
        print("ไม่พบรายชื่อครูในไฟล์", file=sys.stderr)
        return 2
    if args.template.suffix.lower() != ".docx":
        print("template ต้องเป็น .docx", file=sys.stderr)
        return 2

    out_dir = args.out
    raw_dir = out_dir / "raw" if args.save_raw else None
    all_records: list[SubjectRecord] = []
    failed: list[str] = []

    for item in teachers:
        username = item["username"]
        print(f"กำลังดึงข้อมูล {username} ...", flush=True)
        try:
            records = fetch_teacher(username, item["password"], raw_dir)
        except Exception as exc:
            failed.append(f"{username}: {exc}")
            print(f"  ล้มเหลว: {exc}", file=sys.stderr, flush=True)
            continue

        all_records.extend(records)
        docx_name = f"ปถ05_{normalize_filename(username)}_{normalize_filename(records[0].teacher_name if records else username)}.docx"
        fill_docx(args.template, records, out_dir / "docx" / docx_name, args.term, args.year, args.date)
        print(f"  สำเร็จ {len(records)} รายวิชา", flush=True)

    write_summary(all_records, out_dir / "pt05_summary.xlsx")
    write_json(all_records, out_dir / "pt05_summary.json")
    if failed:
        (out_dir / "failed.txt").write_text("\n".join(failed), encoding="utf-8")
        print(f"เสร็จบางส่วน: สำเร็จ {len(all_records)} รายการ, ล้มเหลว {len(failed)} บัญชี", file=sys.stderr)
        return 1

    if args.save_raw and raw_dir and raw_dir.exists():
        try:
            from attendance_risk_tool import analyze_attendance, export_excel, export_html_dashboard, export_json
            print("กำลังวิเคราะห์เวลาเรียนและสังเคราะห์ AI สำหรับนักเรียนกลุ่มเสี่ยง (<60% และ <80%) ...", flush=True)
            profiles, subject_records = analyze_attendance(raw_dir, out_dir / "pt05_summary.json")
            export_excel(profiles, subject_records, out_dir / "student_attendance_risk_summary.xlsx")
            export_html_dashboard(profiles, subject_records, out_dir / "attendance_risk_dashboard.html")
            export_json(profiles, subject_records, out_dir / "student_attendance_risk_summary.json")
            print(f"  วิเคราะห์เวลาเรียนเสร็จสิ้น: พบนักเรียนกลุ่มเสี่ยง {len(profiles)} คน ({out_dir / 'student_attendance_risk_summary.xlsx'})", flush=True)
        except Exception as exc:
            print(f"  คำเตือน: วิเคราะห์เวลาเรียนไม่สำเร็จ: {exc}", file=sys.stderr, flush=True)

    print(f"เสร็จแล้ว: {out_dir}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
